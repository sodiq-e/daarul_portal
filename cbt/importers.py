import csv
import json
import logging
import re
import urllib.request
from dataclasses import dataclass, field
from io import BytesIO, StringIO
from typing import Any, Dict, List, Optional, Tuple, Union

from django.core.files.base import ContentFile
from django.core.files.uploadedfile import UploadedFile
from django.db import transaction
from django.utils.functional import cached_property

from .models import CBTChoice, CBTExam, CBTQuestion, QuestionBank

logger = logging.getLogger(__name__)

FIELD_ALIASES = {
    'prompt': 'prompt',
    'question': 'prompt',
    'question_text': 'prompt',
    'text': 'prompt',
    'question_text': 'prompt',
    'question_text': 'prompt',

    'question_type': 'question_type',
    'type': 'question_type',
    'qtype': 'question_type',

    'marks': 'mark_value',
    'mark': 'mark_value',
    'score': 'mark_value',
    'points': 'mark_value',
    'weight': 'mark_value',

    'difficulty': 'difficulty',
    'level': 'difficulty',

    'topic': 'topic',
    'category': 'topic',
    'subject': 'topic',

    'explanation': 'explanation',
    'solution': 'explanation',
    'working': 'explanation',
    'answer_explanation': 'explanation',

    'image': 'image_url',
    'image_url': 'image_url',
    'imageurl': 'image_url',

    'tags': 'tags',
    'tag': 'tags',
    'keywords': 'tags',
    'labels': 'tags',

    'choices': 'choices',
    'options': 'choices',

    'correct_answer': 'correct_answer',
    'correct_answers': 'correct_answer',
    'answer': 'correct_answer',
    'correct': 'correct_answer',
    'correct_option': 'correct_answer',
    'correct_options': 'correct_answer',

    'order': 'order',
    'position': 'order',
    'index': 'order',

    'question_bank': 'question_bank',
    'bank': 'question_bank',
    'question_bank_name': 'question_bank',
}

CHOICE_COLUMN_PATTERN = re.compile(r'^(?:option|choice|answer|opt|ans|choice_)?[ _-]?([a-e1-5])$', re.I)
QUESTION_TYPE_ALIASES = {
    'mcq': CBTQuestion.MCQ,
    'multiple_choice': CBTQuestion.MCQ,
    'multiple choice': CBTQuestion.MCQ,
    'multiple-choice': CBTQuestion.MCQ,
    'multiple': CBTQuestion.MULTIPLE,
    'multiple_select': CBTQuestion.MULTIPLE,
    'multiple select': CBTQuestion.MULTIPLE,
    'multiple-select': CBTQuestion.MULTIPLE,
    'multiselect': CBTQuestion.MULTIPLE,
    'true_false': CBTQuestion.TRUE_FALSE,
    'true false': CBTQuestion.TRUE_FALSE,
    'true/false': CBTQuestion.TRUE_FALSE,
    'truefalse': CBTQuestion.TRUE_FALSE,
    'boolean': CBTQuestion.TRUE_FALSE,
    'tf': CBTQuestion.TRUE_FALSE,
    'short_answer': CBTQuestion.SHORT_ANSWER,
    'short answer': CBTQuestion.SHORT_ANSWER,
    'short-answer': CBTQuestion.SHORT_ANSWER,
    'essay': CBTQuestion.SHORT_ANSWER,
    'long_answer': CBTQuestion.SHORT_ANSWER,
    'long answer': CBTQuestion.SHORT_ANSWER,
}
DIFFICULTY_ALIASES = {
    'easy': CBTQuestion.DIFFICULTY_EASY,
    'e': CBTQuestion.DIFFICULTY_EASY,
    'medium': CBTQuestion.DIFFICULTY_MEDIUM,
    'm': CBTQuestion.DIFFICULTY_MEDIUM,
    'normal': CBTQuestion.DIFFICULTY_MEDIUM,
    'hard': CBTQuestion.DIFFICULTY_HARD,
    'h': CBTQuestion.DIFFICULTY_HARD,
}
TRUE_VALUES = {'true', 't', 'yes', 'y', '1'}
FALSE_VALUES = {'false', 'f', 'no', 'n', '0'}
LETTER_KEYS = {'a': 0, 'b': 1, 'c': 2, 'd': 3, 'e': 4}
NUMBER_KEYS = {'1': 0, '2': 1, '3': 2, '4': 3, '5': 4}

MAX_CHOICES = 5

@dataclass
class QuestionImportRow:
    row_number: int
    original: Dict[str, Any]
    normalized: Dict[str, Any] = field(default_factory=dict)
    errors: List[str] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    valid: bool = False

    @cached_property
    def preview(self) -> Dict[str, Any]:
        return {
            'row_number': self.row_number,
            'fields': self.normalized,
            'errors': self.errors,
            'warnings': self.warnings,
            'valid': self.valid,
        }


class QuestionImportError(Exception):
    pass


class QuestionImporter:
    def __init__(self, user=None, target_exam: Optional[CBTExam] = None, target_bank: Optional[QuestionBank] = None):
        self.user = user
        self.target_exam = target_exam
        self.target_bank = target_bank

    @staticmethod
    def _normalize_key(key: str) -> str:
        if not isinstance(key, str):
            return ''
        normalized = key.strip().lower().replace(' ', '_').replace('-', '_')
        return FIELD_ALIASES.get(normalized, normalized)

    @staticmethod
    def _normalize_text(value: Any) -> str:
        if value is None:
            return ''
        if isinstance(value, str):
            return value.strip()
        return str(value).strip()

    @staticmethod
    def _split_values(value: str) -> List[str]:
        if value is None:
            return []
        if isinstance(value, list):
            return [str(v).strip() for v in value if str(v).strip()]
        text = str(value).strip()
        if not text:
            return []
        return [item.strip() for item in re.split(r'[;|\n\r]+', text) if item.strip()]

    @staticmethod
    def _parse_float(value: Any, default: float = 1.0) -> float:
        try:
            return float(str(value).strip()) if value is not None and str(value).strip() else default
        except (TypeError, ValueError):
            return default

    @staticmethod
    def _normalize_question_type(value: Any) -> str:
        if not value:
            return ''
        candidate = str(value).strip().lower().replace(' ', '_').replace('-', '_')
        return QUESTION_TYPE_ALIASES.get(candidate, candidate)

    @staticmethod
    def _normalize_difficulty(value: Any) -> str:
        if not value:
            return ''
        candidate = str(value).strip().lower()
        return DIFFICULTY_ALIASES.get(candidate, candidate)

    @staticmethod
    def _extract_choice_columns(record: Dict[str, Any]) -> List[Dict[str, Any]]:
        choices: List[Dict[str, Any]] = []
        for key, value in record.items():
            match = CHOICE_COLUMN_PATTERN.match(key)
            if not match:
                continue
            label = match.group(1).lower()
            if label in LETTER_KEYS:
                index = LETTER_KEYS[label]
            elif label in NUMBER_KEYS:
                index = NUMBER_KEYS[label]
            else:
                continue
            text = str(value).strip()
            if text:
                choices.append({'index': index, 'text': text, 'is_correct': False})
        choices.sort(key=lambda item: item['index'])
        return [{'text': item['text'], 'is_correct': item['is_correct']} for item in choices]

    @staticmethod
    def _build_choice_labels(choices: List[Dict[str, Any]]) -> List[str]:
        labels = []
        for idx, choice in enumerate(choices):
            label = chr(ord('A') + idx) if idx < 26 else str(idx + 1)
            labels.append(label)
        return labels

    @staticmethod
    def _normalize_correct_answers(value: Any) -> List[str]:
        if isinstance(value, list):
            return [str(item).strip() for item in value if str(item).strip()]
        return QuestionImporter._split_values(str(value))

    @staticmethod
    def _match_correct_answer(answer: str, choice_text: str, index: int, label: str) -> bool:
        answer_lower = answer.strip().lower()
        return (
            answer_lower == str(index + 1) or
            answer_lower == label.lower() or
            answer_lower == choice_text.strip().lower()
        )

    @staticmethod
    def _is_true_false_choice_list(choices: List[str]) -> bool:
        normalized = [choice.strip().lower() for choice in choices if choice.strip()]
        return set(normalized) <= {'true', 'false', 't', 'f', 'yes', 'no', 'y', 'n'} and len(normalized) == 2

    def _merge_choice_columns(self, record: Dict[str, Any], normalized: Dict[str, Any]) -> None:
        choice_columns = self._extract_choice_columns(record)
        if choice_columns:
            normalized['choices'] = choice_columns

    def _extract_choices(self, record: Dict[str, Any], normalized: Dict[str, Any]) -> None:
        if 'choices' in record and record['choices'] not in (None, ''):
            raw_choices = record['choices']
            if isinstance(raw_choices, list):
                normalized['choices'] = [
                    {'text': self._normalize_text(item.get('text') if isinstance(item, dict) else item),
                     'is_correct': bool(item.get('is_correct')) if isinstance(item, dict) else False}
                    for item in raw_choices
                ]
            else:
                normalized['choices'] = [
                    {'text': self._normalize_text(value), 'is_correct': False}
                    for value in self._split_values(str(raw_choices))
                ]
            return

        self._merge_choice_columns(record, normalized)

        if normalized.get('question_type') == CBTQuestion.TRUE_FALSE and not normalized.get('choices'):
            normalized['choices'] = [
                {'text': 'True', 'is_correct': False},
                {'text': 'False', 'is_correct': False},
            ]

    def _apply_correct_answers(self, normalized: Dict[str, Any]) -> None:
        raw_correct = normalized.get('correct_answer')
        if not raw_correct:
            return
        correct_answers = self._normalize_correct_answers(raw_correct)
        if not correct_answers:
            return
        choices = normalized.get('choices', [])
        labels = self._build_choice_labels(choices)
        selected = set()
        for raw in correct_answers:
            for idx, choice in enumerate(choices):
                if self._match_correct_answer(raw, choice['text'], idx, labels[idx]):
                    selected.add(idx)
        if selected:
            for idx, choice in enumerate(choices):
                choice['is_correct'] = idx in selected
        elif normalized.get('question_type') == CBTQuestion.TRUE_FALSE and len(choices) == 2:
            answer = correct_answers[0].strip().lower()
            if answer in TRUE_VALUES:
                choices[0]['is_correct'] = True
                choices[1]['is_correct'] = False
            elif answer in FALSE_VALUES:
                choices[0]['is_correct'] = False
                choices[1]['is_correct'] = True

    def _detect_question_type(self, normalized: Dict[str, Any]) -> str:
        explicit = normalized.get('question_type', '').lower()
        if explicit in QUESTION_TYPE_ALIASES.values():
            return explicit
        if explicit:
            candidate = QUESTION_TYPE_ALIASES.get(explicit, '')
            if candidate:
                return candidate
        choices = normalized.get('choices', [])
        correct_count = sum(1 for choice in choices if choice.get('is_correct'))
        texts = [self._normalize_text(choice.get('text')) for choice in choices if choice.get('text')]
        if normalized.get('question_type') in QUESTION_TYPE_ALIASES.values():
            return normalized['question_type']
        if len(choices) == 0:
            return CBTQuestion.SHORT_ANSWER
        if self._is_true_false_choice_list(texts):
            return CBTQuestion.TRUE_FALSE
        if correct_count > 1:
            return CBTQuestion.MULTIPLE
        if correct_count == 0 and len(choices) == 2 and self._is_true_false_choice_list(texts):
            return CBTQuestion.TRUE_FALSE
        if len(choices) >= 2:
            return CBTQuestion.MCQ
        return CBTQuestion.MCQ

    def _map_choice_correctness(self, normalized: Dict[str, Any]) -> None:
        if normalized.get('question_type') == CBTQuestion.SHORT_ANSWER:
            normalized['choices'] = []
            return
        self._apply_correct_answers(normalized)
        choices = normalized.get('choices', [])
        if normalized.get('question_type') == CBTQuestion.TRUE_FALSE and len(choices) == 2 and not any(choice['is_correct'] for choice in choices):
            choices[0]['is_correct'] = True

    def _normalize_row(self, raw_record: Dict[str, Any], row_number: int) -> QuestionImportRow:
        record: Dict[str, Any] = {}
        for raw_key, raw_value in raw_record.items():
            key = self._normalize_key(raw_key)
            if not key:
                continue
            record[key] = raw_value

        normalized: Dict[str, Any] = {
            'prompt': self._normalize_text(record.get('prompt')),
            'question_type': self._normalize_question_type(record.get('question_type')),
            'mark_value': self._parse_float(record.get('mark_value'), 1.0),
            'difficulty': self._normalize_difficulty(record.get('difficulty')) or CBTQuestion.DIFFICULTY_MEDIUM,
            'topic': self._normalize_text(record.get('topic')),
            'explanation': self._normalize_text(record.get('explanation')),
            'tags': self._normalize_text(record.get('tags')),
            'image_url': self._normalize_text(record.get('image_url')),
            'order': int(self._parse_float(record.get('order'), row_number - 1)),
            'question_bank': self._normalize_text(record.get('question_bank')),
            'correct_answer': self._normalize_text(record.get('correct_answer')),
            'choices': [],
        }

        self._extract_choices(record, normalized)
        normalized['question_type'] = self._detect_question_type(normalized)
        self._map_choice_correctness(normalized)

        row = QuestionImportRow(row_number=row_number, original=raw_record, normalized=normalized)
        self.validate_row(row)
        return row

    def validate_row(self, row: QuestionImportRow) -> None:
        self._validate_row(row)

    def _validate_row(self, row: QuestionImportRow) -> None:
        normalized = row.normalized
        if not normalized['prompt']:
            row.errors.append('Missing question prompt.')
        if normalized['mark_value'] <= 0:
            row.errors.append('Marks must be a positive number.')
        if normalized['difficulty'] not in dict(CBTQuestion.DIFFICULTY_CHOICES):
            row.warnings.append('Unknown difficulty, defaulting to medium.')
            normalized['difficulty'] = CBTQuestion.DIFFICULTY_MEDIUM
        if normalized['question_type'] not in dict(CBTQuestion.QUESTION_TYPE_CHOICES):
            row.errors.append(f"Unsupported question type '{normalized['question_type']}'.")
        if normalized['question_type'] in (CBTQuestion.MCQ, CBTQuestion.MULTIPLE, CBTQuestion.TRUE_FALSE):
            if not normalized['choices']:
                row.errors.append('Missing answer choices for this question type.')
            else:
                texts = [self._normalize_text(choice.get('text')) for choice in normalized['choices']]
                if len(set(texts)) != len(texts):
                    row.errors.append('Duplicate answer choices were detected.')
                if not any(choice.get('is_correct') for choice in normalized['choices']):
                    row.errors.append('Missing correct answer choice.')
                if normalized['question_type'] == CBTQuestion.TRUE_FALSE and len(normalized['choices']) != 2:
                    row.warnings.append('True/False questions should have exactly two choices.')
        row.valid = not row.errors

    def _validate_duplicates(self, rows: List[QuestionImportRow]) -> None:
        seen: Dict[str, List[int]] = {}
        for row in rows:
            prompt = row.normalized.get('prompt', '').strip().lower()
            if not prompt:
                continue
            seen.setdefault(prompt, []).append(row.row_number)
        for prompt, row_numbers in seen.items():
            if len(row_numbers) > 1:
                message = f'Duplicate question prompt detected in rows: {", ".join(map(str, row_numbers))}. '
                for row in rows:
                    if row.row_number in row_numbers:
                        row.errors.append(message)
                        row.valid = False

    def parse_file(self, upload: UploadedFile) -> List[QuestionImportRow]:
        filename = upload.name or ''
        ext = filename.split('.')[-1].lower()
        if ext == 'json':
            records = self._parse_json(upload)
        elif ext == 'csv':
            records = self._parse_csv(upload)
        elif ext in ('xlsx', 'xls'):
            records = self._parse_xlsx(upload)
        elif ext == 'docx':
            records = self._parse_docx(upload)
        else:
            records = self._parse_text(upload)

        rows = [self._normalize_row(record, idx + 1) for idx, record in enumerate(records)]
        self._validate_duplicates(rows)
        return rows

    def _parse_json(self, upload: UploadedFile) -> List[Dict[str, Any]]:
        try:
            payload = json.loads(upload.read().decode('utf-8'))
        except json.JSONDecodeError as exc:
            raise QuestionImportError(f'JSON file parse error: {exc}')

        if isinstance(payload, dict):
            if 'questions' in payload and isinstance(payload['questions'], list):
                payload = payload['questions']
            elif 'items' in payload and isinstance(payload['items'], list):
                payload = payload['items']
            else:
                payload = [payload]

        if not isinstance(payload, list):
            raise QuestionImportError('JSON file must contain an array of questions or a single question object.')
        return [self._flatten_json_question(entry) for entry in payload]

    def _flatten_json_question(self, entry: Any) -> Dict[str, Any]:
        if not isinstance(entry, dict):
            raise QuestionImportError('Each JSON question item must be an object.')
        return entry

    def _parse_csv(self, upload: UploadedFile) -> List[Dict[str, Any]]:
        text = upload.read().decode('utf-8-sig')
        reader = csv.DictReader(StringIO(text))
        if not reader.fieldnames:
            raise QuestionImportError('CSV file must contain a header row.')
        records = []
        for row in reader:
            records.append({key: value for key, value in row.items()})
        return records

    def _parse_text(self, upload: UploadedFile) -> List[Dict[str, Any]]:
        text = upload.read().decode('utf-8')
        json_candidate = self._extract_json_from_text(text)
        if json_candidate is not None:
            return [self._flatten_json_question(entry) for entry in json_candidate]
        return self._parse_plain_text(text)

    def _parse_docx(self, upload: UploadedFile) -> List[Dict[str, Any]]:
        try:
            from docx import Document as DocxDocument
        except ImportError:
            raise QuestionImportError('DOCX import requires python-docx.')
        document = DocxDocument(upload)
        text = '\n\n'.join([para.text for para in document.paragraphs if para.text.strip()])
        if not text:
            raise QuestionImportError('Word document is empty.')
        json_candidate = self._extract_json_from_text(text)
        if json_candidate is not None:
            return [self._flatten_json_question(entry) for entry in json_candidate]
        return self._parse_plain_text(text)

    def _parse_xlsx(self, upload: UploadedFile) -> List[Dict[str, Any]]:
        try:
            import openpyxl
        except ImportError:
            raise QuestionImportError('Excel import requires openpyxl.')
        data = upload.read()
        workbook = openpyxl.load_workbook(filename=BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise QuestionImportError('Excel workbook is empty.')
        headers = [self._normalize_text(cell).lower() if cell is not None else '' for cell in rows[0]]
        records: List[Dict[str, Any]] = []
        for row in rows[1:]:
            record: Dict[str, Any] = {}
            for idx, value in enumerate(row or []):
                if idx >= len(headers):
                    continue
                record[headers[idx]] = value
            records.append(record)
        return records

    def _extract_json_from_text(self, text: str) -> Optional[List[Dict[str, Any]]]:
        cleaned = text.strip()
        if not cleaned:
            return None
        if cleaned.startswith('{') or cleaned.startswith('['):
            try:
                data = json.loads(cleaned)
                if isinstance(data, list):
                    return data
                if isinstance(data, dict) and 'questions' in data and isinstance(data['questions'], list):
                    return data['questions']
            except json.JSONDecodeError:
                pass
        match = re.search(r'([\[{].*[\]}])', text, re.S)
        if match:
            try:
                data = json.loads(match.group(1))
                if isinstance(data, list):
                    return data
                if isinstance(data, dict) and 'questions' in data and isinstance(data['questions'], list):
                    return data['questions']
            except json.JSONDecodeError:
                pass
        return None

    def _parse_plain_text(self, text: str) -> List[Dict[str, Any]]:
        questions: List[Dict[str, Any]] = []
        sections = [section.strip() for section in re.split(r'\n\s*\n+', text.strip()) if section.strip()]
        for section in sections:
            lines = [line.strip() for line in section.splitlines() if line.strip()]
            if not lines:
                continue
            data: Dict[str, Any] = {
                'prompt': '',
                'question_type': '',
                'mark_value': 1.0,
                'difficulty': '',
                'topic': '',
                'explanation': '',
                'tags': '',
                'image_url': '',
                'correct_answer': '',
                'choices': [],
                'question_bank': '',
                'order': 0,
            }
            current_key = None
            for line in lines:
                lower = line.lower()
                if lower.startswith('prompt:') or lower.startswith('question:'):
                    data['prompt'] = line.split(':', 1)[1].strip()
                    current_key = 'prompt'
                    continue
                if lower.startswith('type:'):
                    data['question_type'] = line.split(':', 1)[1].strip()
                    current_key = 'question_type'
                    continue
                if lower.startswith('marks:') or lower.startswith('mark:') or lower.startswith('score:') or lower.startswith('points:'):
                    data['mark_value'] = line.split(':', 1)[1].strip()
                    current_key = 'mark_value'
                    continue
                if lower.startswith('difficulty:') or lower.startswith('level:'):
                    data['difficulty'] = line.split(':', 1)[1].strip()
                    current_key = 'difficulty'
                    continue
                if lower.startswith('topic:') or lower.startswith('category:'):
                    data['topic'] = line.split(':', 1)[1].strip()
                    current_key = 'topic'
                    continue
                if lower.startswith('explanation:') or lower.startswith('solution:') or lower.startswith('working:'):
                    data['explanation'] = line.split(':', 1)[1].strip()
                    current_key = 'explanation'
                    continue
                if lower.startswith('tags:'):
                    data['tags'] = line.split(':', 1)[1].strip()
                    current_key = 'tags'
                    continue
                if lower.startswith('image:') or lower.startswith('image_url:'):
                    data['image_url'] = line.split(':', 1)[1].strip()
                    current_key = 'image_url'
                    continue
                if lower.startswith('question_bank:') or lower.startswith('bank:'):
                    data['question_bank'] = line.split(':', 1)[1].strip()
                    current_key = 'question_bank'
                    continue
                if lower.startswith('correct_answer:') or lower.startswith('correct answers:') or lower.startswith('answer:') or lower.startswith('correct:'):
                    data['correct_answer'] = line.split(':', 1)[1].strip()
                    current_key = 'correct_answer'
                    continue
                if lower.startswith('choices:') or lower.startswith('options:'):
                    current_key = 'choices'
                    continue
                if current_key == 'choices' and (line.startswith('-') or line.startswith('*')):
                    content = line[1:].strip()
                    is_correct = content.startswith('*') or content.startswith('✓')
                    if is_correct:
                        content = content[1:].strip()
                    data['choices'].append({'text': content, 'is_correct': is_correct})
                    continue
                if current_key == 'choices' and line:
                    data['choices'].append({'text': line, 'is_correct': False})
                    continue
            questions.append(data)
        return questions

    def _row_to_preview(self, row: QuestionImportRow) -> Dict[str, Any]:
        preview = {
            'row_number': row.row_number,
            'prompt': row.normalized.get('prompt', ''),
            'question_type': row.normalized.get('question_type', ''),
            'mark_value': row.normalized.get('mark_value', 1),
            'difficulty': row.normalized.get('difficulty', ''),
            'topic': row.normalized.get('topic', ''),
            'explanation': row.normalized.get('explanation', ''),
            'tags': row.normalized.get('tags', ''),
            'image_url': row.normalized.get('image_url', ''),
            'choices': [choice.get('text', '') for choice in row.normalized.get('choices', [])],
            'correct_answer': '; '.join([
                chr(ord('A') + idx) for idx, choice in enumerate(row.normalized.get('choices', [])) if choice.get('is_correct')
            ]),
            'question_bank': row.normalized.get('question_bank', ''),
            'order': row.normalized.get('order', 0),
            'errors': row.errors,
            'warnings': row.warnings,
            'valid': row.valid,
        }
        return preview

    def build_preview(self, upload: UploadedFile) -> Dict[str, Any]:
        rows = self.parse_file(upload)
        return {
            'row_count': len(rows),
            'valid_count': sum(1 for row in rows if row.valid),
            'invalid_count': sum(1 for row in rows if not row.valid),
            'rows': [self._row_to_preview(row) for row in rows],
        }

    def _resolve_target_bank(self, row: QuestionImportRow) -> Optional[QuestionBank]:
        bank_name = row.normalized.get('question_bank')
        if not bank_name or not self.user:
            return None
        bank_name = bank_name.strip()
        if bank_name.isdigit():
            try:
                return QuestionBank.objects.get(pk=int(bank_name), created_by=self.user)
            except QuestionBank.DoesNotExist:
                return None
        return QuestionBank.objects.filter(name__iexact=bank_name, created_by=self.user).first()

    def save(self, rows: List[Dict[str, Any]], import_target: str = 'exam', bank: Optional[QuestionBank] = None) -> Dict[str, Any]:
        if import_target not in ('exam', 'bank', 'both'):
            raise QuestionImportError('Invalid import target.')
        if import_target in ('bank', 'both') and bank is None:
            raise QuestionImportError('A question bank must be selected when importing into a bank.')
        created = 0
        skipped = 0
        invalid = 0
        errors: List[str] = []
        saved_ids: List[int] = []

        with transaction.atomic():
            for idx, row in enumerate(rows, start=1):
                if isinstance(row, dict) and 'errors' in row and row['errors']:
                    invalid += 1
                    continue
                normalized = row if isinstance(row, dict) else row.get('normalized', {})
                if not normalized.get('prompt'):
                    invalid += 1
                    continue
                if import_target in ('exam', 'both') and self.target_exam is None:
                    raise QuestionImportError('Target exam is required for exam import.')
                bank_target = bank if import_target in ('bank', 'both') else None
                if bank_target is None:
                    if normalized.get('question_bank'):
                        bank_target = self._resolve_target_bank(QuestionImportRow(row_number=idx, original={}, normalized=normalized))
                if import_target == 'exam':
                    target_exam = self.target_exam
                else:
                    target_exam = self.target_exam if import_target == 'both' else None
                if normalized['question_type'] == CBTQuestion.SHORT_ANSWER:
                    normalized['choices'] = []
                question = CBTQuestion.objects.create(
                    exam=target_exam,
                    question_bank=bank_target,
                    prompt=normalized['prompt'],
                    question_type=normalized['question_type'],
                    mark_value=normalized.get('mark_value', 1.0),
                    explanation=normalized.get('explanation', ''),
                    topic=normalized.get('topic', ''),
                    difficulty=normalized.get('difficulty', CBTQuestion.DIFFICULTY_MEDIUM),
                    tags=normalized.get('tags', ''),
                    order=normalized.get('order', 0),
                    is_active=True,
                )
                if normalized.get('image_url'):
                    self._attach_image(question, normalized['image_url'])
                for choice_idx, choice in enumerate(normalized.get('choices', [])):
                    text = self._normalize_text(choice.get('text'))
                    if not text:
                        continue
                    CBTChoice.objects.create(
                        question=question,
                        text=text,
                        is_correct=bool(choice.get('is_correct')),
                        order=choice_idx,
                    )
                created += 1
                saved_ids.append(question.id)

        return {
            'created': created,
            'invalid': invalid,
            'skipped': skipped,
            'errors': errors,
            'saved_ids': saved_ids,
        }

    def _attach_image(self, question: CBTQuestion, image_url: str) -> None:
        image_url = image_url.strip()
        if not image_url:
            return
        if image_url.startswith(('http://', 'https://')):
            try:
                with urllib.request.urlopen(image_url, timeout=10) as response:
                    content = response.read()
                    filename = image_url.split('/')[-1].split('?')[0] or 'image.jpg'
                    question.image.save(filename, ContentFile(content), save=True)
            except Exception as exc:
                logger.warning('Unable to import question image from URL %s: %s', image_url, exc)
        else:
            logger.warning('Image URL is not remote, skipping import for %s', image_url)


class ImportPreviewSerializer:
    @staticmethod
    def serialize_row(row: QuestionImportRow) -> Dict[str, Any]:
        return {
            'row_number': row.row_number,
            'prompt': row.normalized.get('prompt', ''),
            'question_type': row.normalized.get('question_type', ''),
            'mark_value': row.normalized.get('mark_value', 1.0),
            'difficulty': row.normalized.get('difficulty', ''),
            'topic': row.normalized.get('topic', ''),
            'explanation': row.normalized.get('explanation', ''),
            'tags': row.normalized.get('tags', ''),
            'image_url': row.normalized.get('image_url', ''),
            'question_bank': row.normalized.get('question_bank', ''),
            'order': row.normalized.get('order', 0),
            'choices': [choice.get('text', '') for choice in row.normalized.get('choices', [])],
            'correct_answer': '; '.join([
                chr(ord('A') + idx) for idx, choice in enumerate(row.normalized.get('choices', [])) if choice.get('is_correct')
            ]),
            'errors': row.errors,
            'warnings': row.warnings,
            'valid': row.valid,
        }

    @staticmethod
    def serialize(rows: List[QuestionImportRow]) -> Dict[str, Any]:
        return {
            'row_count': len(rows),
            'valid_count': sum(1 for row in rows if row.valid),
            'invalid_count': sum(1 for row in rows if not row.valid),
            'rows': [ImportPreviewSerializer.serialize_row(row) for row in rows],
        }
