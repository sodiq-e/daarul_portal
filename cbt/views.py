import csv
import json
import logging
import os
import re
import time
import zipfile
from io import BytesIO, StringIO
from urllib.parse import quote

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from django.core.cache import cache
from django.db import transaction
from django.db.models import Avg, Q, Count, Max
from django.http import FileResponse, HttpResponse, HttpResponseForbidden, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.template.loader import render_to_string
from django.urls import reverse_lazy, reverse
from django.utils import timezone
from django.utils.text import slugify
from django.utils.decorators import method_decorator
from django.views.decorators.csrf import csrf_exempt
from django.views.generic import CreateView, DetailView, ListView, TemplateView, UpdateView
from docx import Document

from .ai_provider import generate_ai_questions
from .forms import CBTExamForm, CBTQuestionForm, CBTChoiceFormSet
from .gemini_service import (
    GeminiAPIError,
    GeminiInvalidAPIKeyError,
    GeminiJSONError,
    GeminiModelNotFoundError,
    GeminiQuotaError,
    GeminiTimeoutError,
    generate_ss1_questions,
    validate_generated_question_payload,
)
from .models import (
    AIRequestMetric,
    CBTAnswer,
    CBTAttemptIntegrityEvent,
    CBTExam,
    CBTChoice,
    CBTQuestion,
    CBTStudentAttempt,
    QuestionBank,
    StudentAttemptQuestion,
)
from .importers import QuestionImporter, ImportPreviewSerializer, QuestionImportError
from .services import build_attempt_context, create_attempt, grade_attempt, save_answer

AI_GENERATION_THROTTLE_SECONDS = int(os.getenv('AI_GENERATION_THROTTLE_SECONDS', '8'))
AI_GENERATION_SESSION_KEY = 'cbt_ai_generation'
AI_GENERATION_RATE_LIMIT_PER_MINUTE = int(os.getenv('AI_GENERATION_RATE_LIMIT_PER_MINUTE', '5'))


def _increment_ai_generation_failure(request):
    throttle_data = request.session.get(AI_GENERATION_SESSION_KEY, {})
    throttle_data['daily_failed_requests'] = throttle_data.get('daily_failed_requests', 0) + 1
    request.session[AI_GENERATION_SESSION_KEY] = throttle_data
    request.session.modified = True


def _enforce_ai_rate_limit(request):
    if not getattr(settings, 'AI_GENERATION_ENABLED', True):
        return JsonResponse({'error': 'AI question generation is currently disabled. Please try again later.'}, status=503)

    rate_limit_key = f'cbt:ai:rate-limit:{request.user.pk}'
    current = cache.get(rate_limit_key) or 0
    if current >= AI_GENERATION_RATE_LIMIT_PER_MINUTE:
        return JsonResponse({'error': 'AI question generation rate limit exceeded. Please wait a minute and try again.'}, status=429)

    cache.set(rate_limit_key, current + 1, 60)
    return None


def is_cbt_teacher(user):
    try:
        return user.profile.is_approved and user.groups.filter(name__in=['Teacher', 'Staff']).exists()
    except Exception:
        return user.is_staff


def is_cbt_authenticated(user):
    try:
        return user.profile.is_approved
    except Exception:
        return user.is_authenticated


def is_cbt_student(user):
    try:
        return user.profile.is_approved and hasattr(user, 'student_profile')
    except Exception:
        return False


@method_decorator(login_required, name='dispatch')
class CBTExamListView(LoginRequiredMixin, UserPassesTestMixin, ListView):
    model = CBTExam
    template_name = 'cbt/exam_list.html'
    context_object_name = 'exams'
    paginate_by = 20

    def test_func(self):
        return is_cbt_teacher(self.request.user)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['exam_mode_choices'] = CBTExam.EXAM_MODE_CHOICES
        return context


@method_decorator(login_required, name='dispatch')
class CBTExamCreateView(LoginRequiredMixin, UserPassesTestMixin, CreateView):
    model = CBTExam
    form_class = CBTExamForm
    template_name = 'cbt/exam_form.html'
    success_url = reverse_lazy('cbt:exam_list')

    def test_func(self):
        return is_cbt_teacher(self.request.user)

    def form_valid(self, form):
        form.instance.created_by = self.request.user
        messages.success(self.request, 'CBT exam created successfully.')
        return super().form_valid(form)


@method_decorator(login_required, name='dispatch')
class CBTExamUpdateView(LoginRequiredMixin, UserPassesTestMixin, UpdateView):
    model = CBTExam
    form_class = CBTExamForm
    template_name = 'cbt/exam_form.html'
    success_url = reverse_lazy('cbt:exam_list')

    def test_func(self):
        return is_cbt_teacher(self.request.user)

    def form_valid(self, form):
        messages.success(self.request, 'CBT exam updated successfully.')
        return super().form_valid(form)


@method_decorator(login_required, name='dispatch')
class CBTQuestionListView(LoginRequiredMixin, UserPassesTestMixin, TemplateView):
    template_name = 'cbt/question_list.html'

    def test_func(self):
        return is_cbt_teacher(self.request.user)

    def get_context_data(self, **kwargs):
        exam = get_object_or_404(CBTExam, pk=self.kwargs.get('exam_pk'))
        return {
            'exam': exam,
            'questions': exam.questions.all(),
        }


@method_decorator(login_required, name='dispatch')
class CBTQuestionCreateView(LoginRequiredMixin, UserPassesTestMixin, TemplateView):
    template_name = 'cbt/question_form.html'

    def test_func(self):
        return is_cbt_teacher(self.request.user)

    def get(self, request, *args, **kwargs):
        exam = get_object_or_404(CBTExam, pk=self.kwargs.get('exam_pk'))
        form = CBTQuestionForm()
        formset = CBTChoiceFormSet()
        return render(request, self.template_name, {'exam': exam, 'form': form, 'formset': formset})

    def post(self, request, *args, **kwargs):
        exam = get_object_or_404(CBTExam, pk=self.kwargs.get('exam_pk'))
        form = CBTQuestionForm(request.POST)
        formset = CBTChoiceFormSet(request.POST)
        if form.is_valid() and formset.is_valid():
            question = form.save(commit=False)
            question.exam = exam
            question.save()
            formset.instance = question
            formset.save()
            messages.success(request, 'Question created successfully.')
            return redirect('cbt:question_list', exam_pk=exam.pk)
        return render(request, self.template_name, {'exam': exam, 'form': form, 'formset': formset})


@method_decorator(login_required, name='dispatch')
class CBTQuestionUpdateView(LoginRequiredMixin, UserPassesTestMixin, TemplateView):
    template_name = 'cbt/question_form.html'

    def test_func(self):
        return is_cbt_teacher(self.request.user)

    def get(self, request, *args, **kwargs):
        question = get_object_or_404(CBTQuestion, pk=self.kwargs.get('question_pk'))
        form = CBTQuestionForm(instance=question)
        formset = CBTChoiceFormSet(instance=question)
        return render(request, self.template_name, {'exam': question.exam, 'form': form, 'formset': formset, 'question': question})

    def post(self, request, *args, **kwargs):
        question = get_object_or_404(CBTQuestion, pk=self.kwargs.get('question_pk'))
        form = CBTQuestionForm(request.POST, instance=question)
        formset = CBTChoiceFormSet(request.POST, instance=question)
        if form.is_valid() and formset.is_valid():
            form.save()
            formset.save()
            messages.success(request, 'Question updated successfully.')
            return redirect('cbt:question_list', exam_pk=question.exam.pk)
        return render(request, self.template_name, {'exam': question.exam, 'form': form, 'formset': formset, 'question': question})


class PracticeExamListView(TemplateView):
    template_name = 'cbt/practice_exam_list.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['practice_exams'] = CBTExam.objects.filter(
            exam_mode=CBTExam.PRACTICE,
            is_active=True,
            is_published=True,
        ).order_by('name')
        return context


class RealExamListView(LoginRequiredMixin, ListView):
    model = CBTExam
    template_name = 'cbt/real_exam_list.html'
    context_object_name = 'real_exams'

    def get_queryset(self):
        return CBTExam.objects.filter(
            exam_mode=CBTExam.REAL,
            is_active=True,
            is_published=True,
        ).order_by('name')


@login_required
def real_exam_detail(request, pk):
    exam = get_object_or_404(CBTExam, pk=pk, exam_mode=CBTExam.REAL)
    return render(request, 'cbt/real_exam_detail.html', {'exam': exam})


def _ensure_session(request):
    if not request.session.session_key:
        request.session.create()
    return request.session.session_key


def start_practice_exam(request, pk):
    exam = get_object_or_404(CBTExam, pk=pk, exam_mode=CBTExam.PRACTICE, is_published=True, is_active=True)
    session_key = _ensure_session(request)
    attempt = create_attempt(exam=exam, session_key=session_key)
    return redirect('cbt:attempt_detail', uuid=attempt.uuid)


@login_required
def start_real_exam(request, pk):
    exam = get_object_or_404(CBTExam, pk=pk, exam_mode=CBTExam.REAL, is_published=True, is_active=True)
    attempt = create_attempt(exam=exam, student=request.user)
    return redirect('cbt:attempt_detail', uuid=attempt.uuid)


def attempt_detail(request, uuid):
    attempt = get_object_or_404(CBTStudentAttempt, uuid=uuid)
    # Permission enforcement:
    # - If the attempt belongs to an authenticated student, allow only that student, the exam owner (teacher), or staff.
    # - If the attempt is a practice attempt (no student), allow only when session_key matches or the exam owner/staff.
    if attempt.student:
        # attempt is tied to a user account; require authentication
        if not request.user.is_authenticated:
            return HttpResponseForbidden()
        # allow student themselves, the exam creator (teacher), or staff
        if not (request.user == attempt.student or request.user == attempt.exam.created_by or request.user.is_staff):
            return HttpResponseForbidden()
    else:
        # practice attempt tied to a session_key
        session_key = _ensure_session(request)
        if attempt.session_key and attempt.session_key != session_key:
            # allow exam owner or staff to inspect practice attempts
            if not (request.user.is_authenticated and (request.user == attempt.exam.created_by or request.user.is_staff)):
                return HttpResponseForbidden()

    if attempt.is_submitted:
        context = build_attempt_context(attempt)
        return render(request, 'cbt/attempt_result.html', context)

    if request.method == 'POST':
        for attempt_question in attempt.attempt_questions.select_related('question').order_by('randomized_position'):
            question = attempt_question.question
            answer_field = request.POST.get(f'question_{question.pk}')
            selected_choice = None
            text_answer = ''
            if question.question_type in [CBTQuestion.MCQ, CBTQuestion.TRUE_FALSE]:
                try:
                    selected_choice = CBTChoice.objects.get(pk=int(answer_field), question=question)
                except (ValueError, CBTChoice.DoesNotExist, TypeError):
                    selected_choice = None
            else:
                text_answer = answer_field or ''
            save_answer(attempt=attempt, question=question, selected_choice=selected_choice, text_answer=text_answer)
        grade_attempt(attempt)
        messages.success(request, 'Exam submitted and graded. Your score has been recorded.')
        return redirect('cbt:attempt_detail', uuid=attempt.uuid)

    context = build_attempt_context(attempt)
    # Build JSON payload for frontend
    # Use attempt.attempt_questions ordering if available
    attempt_questions = attempt.attempt_questions.select_related('question').prefetch_related('question__choices').order_by('randomized_position')
    questions_list = []
    answers = {a.question_id: a for a in attempt.answers.select_related('selected_choice').all()}
    for aq in attempt_questions:
        q = aq.question
        # build choices in randomized order if present
        if aq.randomized_choice_order:
            try:
                order_ids = json.loads(aq.randomized_choice_order)
            except Exception:
                order_ids = []
            choices_qs = list(q.choices.all())
            choices_map = {c.id: c for c in choices_qs}
            choices = [
                {'id': cid, 'text': choices_map[cid].text} for cid in order_ids if cid in choices_map
            ]
            # append any missing choices
            for c in choices_qs:
                if c.id not in order_ids:
                    choices.append({'id': c.id, 'text': c.text})
        else:
            choices = [{'id': c.id, 'text': c.text} for c in q.choices.all().order_by('order')]

        ans = answers.get(q.id)
        answer_payload = None
        if ans:
            try:
                sel_multi = json.loads(ans.selected_choices) if ans.selected_choices else []
            except Exception:
                sel_multi = []
            answer_payload = {
                'selected_choice_id': ans.selected_choice_id,
                'selected_choice_ids': sel_multi,
                'text_answer': ans.text_answer,
            }

        questions_list.append({
            'id': q.id,
            'prompt': q.prompt,
            'question_type': q.question_type,
            'mark_value': float(q.mark_value),
            'topic': getattr(q, 'topic', ''),
            'difficulty': getattr(q, 'difficulty', ''),
            'choices': choices,
        })

    # inject JSON into context for template
    context['questions_json'] = json.dumps(questions_list)
    # build answers array aligned by position
    answers_arr = []
    for aq in attempt_questions:
        qid = aq.question.id
        a = answers.get(qid)
        if a:
            try:
                sel_multi = json.loads(a.selected_choices) if a.selected_choices else []
            except Exception:
                sel_multi = []
            answers_arr.append({
                'question_id': qid,
                'selected_choice_id': a.selected_choice_id,
                'selected_choice_ids': sel_multi,
                'text_answer': a.text_answer,
            })
        else:
            answers_arr.append(None)

    context['answers_json'] = json.dumps(answers_arr)
    context['flagged_questions_json'] = json.dumps(list(attempt_questions.filter(is_flagged=True).values_list('randomized_position', flat=True)))
    context['total_questions'] = attempt.attempt_questions.count() or attempt.exam.questions.filter(is_active=True).count()
    # Determine remaining time for the current attempt
    elapsed_seconds = max(0, int((timezone.now() - attempt.started_at).total_seconds()))
    total_seconds = int(attempt.exam.duration_minutes * 60)
    context['time_left_seconds'] = max(0, total_seconds - elapsed_seconds)
    return render(request, 'cbt/student_attempt_view.html', context)


def api_save_answer(request):
    if request.method != 'POST':
        return JsonResponse({'error': 'POST required'}, status=400)
    try:
        data = json.loads(request.body.decode())
    except Exception:
        return JsonResponse({'error': 'invalid json'}, status=400)

    attempt_uuid = data.get('attempt_uuid')
    question_id = data.get('question_id')
    selected_choice = data.get('selected_choice_id')
    selected_choice_ids = data.get('selected_choice_ids')
    text_answer = data.get('text_answer', '')

    attempt = get_object_or_404(CBTStudentAttempt, uuid=attempt_uuid)
    # permission check
    if attempt.student:
        # student-owned attempt: must be the student, exam owner, or staff
        if attempt.student != request.user and not (request.user.is_authenticated and (request.user == attempt.exam.created_by or request.user.is_staff)):
            return HttpResponseForbidden()
    else:
        # practice attempt tied to a session_key: require matching session or exam owner/staff
        session_key = _ensure_session(request)
        if not attempt.session_key or attempt.session_key != session_key:
            if not (request.user.is_authenticated and (request.user == attempt.exam.created_by or request.user.is_staff)):
                return HttpResponseForbidden()

    if attempt.is_submitted:
        return JsonResponse({'error': 'attempt already submitted'}, status=400)

    # Verify question belongs to this attempt
    attempt_question = get_object_or_404(StudentAttemptQuestion, attempt=attempt, question_id=question_id)
    question = attempt_question.question
    
    # handle answer updates only when answer data is present
    has_answer_data = selected_choice_ids is not None or selected_choice is not None or (text_answer and text_answer.strip());
    if has_answer_data:
        if selected_choice_ids is not None:
            save_answer(attempt=attempt, question=question, selected_choice=selected_choice_ids, text_answer='')
        else:
            sel_choice_obj = None
            if selected_choice:
                try:
                    sel_choice_obj = CBTChoice.objects.get(pk=int(selected_choice), question=question)
                except CBTChoice.DoesNotExist:
                    return JsonResponse({'error': 'invalid choice for this question'}, status=400)
            save_answer(attempt=attempt, question=question, selected_choice=sel_choice_obj, text_answer=text_answer)

    if 'is_flagged' in data:
        is_flagged = bool(data.get('is_flagged'))
        saq = StudentAttemptQuestion.objects.filter(attempt=attempt, question=question).first()
        if saq:
            saq.is_flagged = is_flagged
            saq.save(update_fields=['is_flagged'])

    integrity_events = data.get('integrity_events')
    if isinstance(integrity_events, list) and integrity_events:
        for event in integrity_events:
            if not isinstance(event, dict):
                continue
            reason = event.get('reason') or 'unknown'
            metadata = event.get('metadata') if isinstance(event.get('metadata'), dict) else {}
            CBTAttemptIntegrityEvent.objects.create(
                attempt=attempt,
                reason=reason,
                metadata=metadata
            )

    return JsonResponse({'status': 'saved', 'last_saved': attempt.last_saved_at.isoformat()})


@csrf_exempt
def api_submit_attempt(request):
    if request.method != 'POST':
        return JsonResponse({'error': 'POST required'}, status=400)

    attempt_uuid = None
    body = request.body.decode(errors='ignore')
    if body:
        try:
            data = json.loads(body)
            attempt_uuid = data.get('attempt_uuid')
        except Exception:
            pass

    if not attempt_uuid:
        attempt_uuid = request.POST.get('attempt_uuid') or request.GET.get('attempt_uuid')

    if not attempt_uuid:
        return JsonResponse({'error': 'attempt_uuid required'}, status=400)

    # Enforce same-origin for beacon/fallback submissions
    origin = request.META.get('HTTP_ORIGIN')
    referer = request.META.get('HTTP_REFERER')
    allowed_origin = f"{request.scheme}://{request.get_host()}"
    if origin and origin != allowed_origin:
        return HttpResponseForbidden()
    if not origin and referer and not referer.startswith(allowed_origin):
        return HttpResponseForbidden()

    attempt = get_object_or_404(CBTStudentAttempt, uuid=attempt_uuid)
    # Permission: allow student owner, exam owner, or staff; for practice attempts require matching session or owner/staff
    if attempt.student:
        if not request.user.is_authenticated:
            return HttpResponseForbidden()
        if not (request.user == attempt.student or request.user == attempt.exam.created_by or request.user.is_staff):
            return HttpResponseForbidden()
    else:
        session_key = _ensure_session(request)
        if not attempt.session_key or attempt.session_key != session_key:
            if not (request.user.is_authenticated and (request.user == attempt.exam.created_by or request.user.is_staff)):
                return HttpResponseForbidden()
    if attempt.is_submitted:
        return JsonResponse({'error': 'already submitted'}, status=400)

    # grade and finalize
    grade_attempt(attempt)
    return JsonResponse({'status': 'submitted', 'score': float(attempt.score or 0)})


@login_required
def api_generate_ai_questions(request, exam_pk):
    if request.method != 'POST':
        return JsonResponse({'error': 'POST required'}, status=400)

    rate_limit_response = _enforce_ai_rate_limit(request)
    if rate_limit_response is not None:
        return rate_limit_response

    if not getattr(settings, 'AI_GENERATION_ENABLED', True):
        return JsonResponse({'error': 'AI question generation is currently disabled. Please try again later.'}, status=503)

    exam = get_object_or_404(CBTExam, pk=exam_pk, created_by=request.user)
    if not exam.allow_ai_questions:
        return JsonResponse({'error': 'AI question generation is not enabled for this exam.'}, status=403)

    # Prevent concurrent generation requests for the same user+exam
    lock_key = f'cbt:ai:lock:{request.user.pk}:{exam.pk}'
    got_lock = cache.add(lock_key, 1, AI_GENERATION_THROTTLE_SECONDS)
    if not got_lock:
        AIRequestMetric.objects.create(
            user=request.user,
            exam=exam,
            request_type=AIRequestMetric.REQUEST_TYPE_GENERATE_AI,
            status=AIRequestMetric.STATUS_THROTTLED,
            error_code='concurrent_request'
        )
        return JsonResponse({'error': 'Another generation request is in progress. Please wait a moment.'}, status=429)

    try:
        data = json.loads(request.body.decode())
    except Exception:
        return JsonResponse({'error': 'Invalid JSON payload.'}, status=400)

    now_ts = int(time.time())
    throttle_data = request.session.get(AI_GENERATION_SESSION_KEY, {})
    last_request_ts = throttle_data.get('last_request_at', 0)
    if now_ts - last_request_ts < AI_GENERATION_THROTTLE_SECONDS:
        return JsonResponse({
            'error': 'Please wait a few seconds before generating questions again. Reduce repeated requests and try again shortly.'
        }, status=429)

    throttle_data['last_request_at'] = now_ts
    throttle_data['daily_request_count'] = throttle_data.get('daily_request_count', 0) + 1
    throttle_data.setdefault('daily_successful_requests', 0)
    throttle_data.setdefault('daily_failed_requests', 0)
    request.session[AI_GENERATION_SESSION_KEY] = throttle_data
    request.session.modified = True

    topic = data.get('topic', '').strip()
    difficulty = data.get('difficulty', CBTQuestion.DIFFICULTY_MEDIUM)
    num_questions = data.get('num_questions', 5)
    try:
        num_questions = int(num_questions)
    except (TypeError, ValueError):
        num_questions = 5

    if num_questions < 1 or num_questions > 20:
        return JsonResponse({'error': 'num_questions must be between 1 and 20.'}, status=400)

    if difficulty not in dict(CBTQuestion.DIFFICULTY_CHOICES):
        return JsonResponse({'error': 'Invalid difficulty value.'}, status=400)

    try:
        start_ts = time.time()
        questions = generate_ai_questions(
            exam=exam,
            topic=topic,
            difficulty=difficulty,
            num_questions=num_questions,
        )
        latency_ms = int((time.time() - start_ts) * 1000)
        throttle_data['daily_successful_requests'] = throttle_data.get('daily_successful_requests', 0) + 1
        request.session[AI_GENERATION_SESSION_KEY] = throttle_data
        request.session.modified = True
        AIRequestMetric.objects.create(
            user=request.user,
            exam=exam,
            request_type=AIRequestMetric.REQUEST_TYPE_GENERATE_AI,
            status=AIRequestMetric.STATUS_SUCCESS,
            latency_ms=latency_ms,
        )
    except GeminiQuotaError as exc:
        _increment_ai_generation_failure(request)
        AIRequestMetric.objects.create(
            user=request.user,
            exam=exam,
            request_type=AIRequestMetric.REQUEST_TYPE_GENERATE_AI,
            status=AIRequestMetric.STATUS_QUOTA,
            error_code=exc.__class__.__name__,
        )
        return JsonResponse({
            'error': str(exc),
            'retryable': True,
        }, status=503)
    except GeminiInvalidAPIKeyError as exc:
        _increment_ai_generation_failure(request)
        AIRequestMetric.objects.create(
            user=request.user,
            exam=exam,
            request_type=AIRequestMetric.REQUEST_TYPE_GENERATE_AI,
            status=AIRequestMetric.STATUS_FAILURE,
            error_code=exc.__class__.__name__,
        )
        return JsonResponse({'error': str(exc)}, status=401)
    except GeminiModelNotFoundError as exc:
        _increment_ai_generation_failure(request)
        AIRequestMetric.objects.create(
            user=request.user,
            exam=exam,
            request_type=AIRequestMetric.REQUEST_TYPE_GENERATE_AI,
            status=AIRequestMetric.STATUS_FAILURE,
            error_code=exc.__class__.__name__,
        )
        return JsonResponse({'error': str(exc)}, status=404)
    except GeminiTimeoutError as exc:
        _increment_ai_generation_failure(request)
        AIRequestMetric.objects.create(
            user=request.user,
            exam=exam,
            request_type=AIRequestMetric.REQUEST_TYPE_GENERATE_AI,
            status=AIRequestMetric.STATUS_FAILURE,
            error_code=exc.__class__.__name__,
        )
        return JsonResponse({'error': str(exc), 'retryable': True}, status=503)
    except GeminiJSONError as exc:
        _increment_ai_generation_failure(request)
        AIRequestMetric.objects.create(
            user=request.user,
            exam=exam,
            request_type=AIRequestMetric.REQUEST_TYPE_GENERATE_AI,
            status=AIRequestMetric.STATUS_FAILURE,
            error_code=exc.__class__.__name__,
        )
        return JsonResponse({'error': 'The AI returned an unexpected response format. Please try again or contact support.'}, status=502)
    except GeminiAPIError as exc:
        _increment_ai_generation_failure(request)
        AIRequestMetric.objects.create(
            user=request.user,
            exam=exam,
            request_type=AIRequestMetric.REQUEST_TYPE_GENERATE_AI,
            status=AIRequestMetric.STATUS_FAILURE,
            error_code=exc.__class__.__name__,
        )
        return JsonResponse({'error': str(exc)}, status=500)
    except Exception as exc:
        _increment_ai_generation_failure(request)
        AIRequestMetric.objects.create(
            user=request.user,
            exam=exam,
            request_type=AIRequestMetric.REQUEST_TYPE_GENERATE_AI,
            status=AIRequestMetric.STATUS_FAILURE,
            error_code=exc.__class__.__name__,
        )
        return JsonResponse({'error': 'AI generation failed. Please try again later.'}, status=500)

    finally:
        try:
            cache.delete(lock_key)
        except Exception:
            pass

    validated_questions = []
    for question in questions:
        valid, error = validate_generated_question_payload(question)
        if not valid:
            return JsonResponse({'error': f'AI returned invalid question data: {error}'}, status=500)
        validated_questions.append(question)

    return JsonResponse({'status': 'ok', 'questions': validated_questions})


@login_required
def api_ai_request_metrics(request):
    if request.method != 'GET':
        return JsonResponse({'error': 'GET required'}, status=400)

    metrics = AIRequestMetric.objects.filter(exam__created_by=request.user)
    today = timezone.localdate()
    summary = list(metrics.values('request_type', 'status').annotate(count=Count('id')))
    today_summary = list(metrics.filter(date=today).values('status').annotate(count=Count('id')))
    return JsonResponse({
        'metrics': summary,
        'today': today_summary,
    })


@login_required
def api_generate_ss1_questions(request):
    """
    Test endpoint for generating SS1 questions from Gemini.

    - Validates GEMINI_API_KEY presence
    - Calls `generate_ss1_questions` which performs cleaning, retries and strict validation
    - Returns ONLY the list of question objects on success (safe=False)
    - On error returns a fallback JSON: {"success": false, "error": "message"}

    Note: This endpoint is for testing the AI integration only. Do NOT wire
    frontend UI to this endpoint until SDK, parsing and validation are confirmed.
    """
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'POST required'}, status=400)

    # add a concurrency guard for SS1 generation as well
    lock_key = f'cbt:ai:ss1:lock:{request.user.pk}'
    got_lock = cache.add(lock_key, 1, AI_GENERATION_THROTTLE_SECONDS)
    if not got_lock:
        return JsonResponse({'success': False, 'error': 'Another AI request is in progress. Please wait.'}, status=429)

    try:
        data = json.loads(request.body.decode() or '{}')
    except Exception:
        return JsonResponse({'success': False, 'error': 'Invalid JSON payload.'}, status=400)

    topic = data.get('topic', '').strip()
    num_questions = data.get('num_questions', 5)
    try:
        num_questions = int(num_questions)
    except (TypeError, ValueError):
        num_questions = 5

    if num_questions < 1 or num_questions > 20:
        return JsonResponse({'success': False, 'error': 'num_questions must be between 1 and 20.'}, status=400)

    try:
        questions = generate_ss1_questions(topic=topic, num_questions=num_questions)
    except EnvironmentError as env_err:
        logging.exception('Environment error when generating SS1 questions')
        return JsonResponse({'success': False, 'error': str(env_err)}, status=500)
    except GeminiQuotaError as exc:
        logging.warning('Gemini quota error for SS1 generation: %s', exc)
        return JsonResponse({'success': False, 'error': str(exc)}, status=503)
    except GeminiInvalidAPIKeyError as exc:
        logging.warning('Gemini auth error for SS1 generation: %s', exc)
        return JsonResponse({'success': False, 'error': str(exc)}, status=401)
    except GeminiModelNotFoundError as exc:
        logging.warning('Gemini model error for SS1 generation: %s', exc)
        return JsonResponse({'success': False, 'error': str(exc)}, status=404)
    except GeminiTimeoutError as exc:
        logging.warning('Gemini timeout for SS1 generation: %s', exc)
        return JsonResponse({'success': False, 'error': str(exc)}, status=503)
    except GeminiJSONError as exc:
        logging.warning('Gemini JSON format error for SS1 generation: %s', exc)
        return JsonResponse({'success': False, 'error': 'The AI returned an unexpected response format. Please try again later.'}, status=502)
    except GeminiAPIError as exc:
        logging.warning('Gemini API error for SS1 generation: %s', exc)
        return JsonResponse({'success': False, 'error': str(exc)}, status=500)
    except Exception as exc:
        logging.exception('Failed to generate SS1 questions')
        return JsonResponse({'success': False, 'error': 'AI generation failed: ' + str(exc)}, status=500)
    finally:
        try:
            cache.delete(lock_key)
        except Exception:
            pass

    # Return the validated list directly (safe=False allows non-dict top-level JSON)
    return JsonResponse(questions, safe=False)


@login_required
def api_save_generated_ai_questions(request, exam_pk):
    if request.method != 'POST':
        return JsonResponse({'error': 'POST required'}, status=400)

    exam = get_object_or_404(CBTExam, pk=exam_pk, created_by=request.user)
    if not exam.allow_ai_questions:
        return JsonResponse({'error': 'AI question generation is not enabled for this exam.'}, status=403)

    try:
        data = json.loads(request.body.decode())
    except Exception:
        return JsonResponse({'error': 'Invalid JSON payload.'}, status=400)

    questions = data.get('questions')
    if not isinstance(questions, list) or not questions:
        return JsonResponse({'error': 'questions must be a non-empty list.'}, status=400)

    created_question_ids = []
    try:
        with transaction.atomic():
            current_max_order = exam.questions.aggregate(max_order=Max('order'))['max_order']
            if current_max_order is None:
                current_max_order = -1

            for idx, question in enumerate(questions):
                valid, error = validate_generated_question_payload(question)
                if not valid:
                    AIRequestMetric.objects.create(
                        user=request.user,
                        exam=exam,
                        request_type=AIRequestMetric.REQUEST_TYPE_GENERATE_AI,
                        status=AIRequestMetric.STATUS_FAILURE,
                        error_code=f'validation_error_{idx}'
                    )
                    return JsonResponse({'error': f'Invalid question payload at index {idx}: {error}'}, status=400)

                qobj = CBTQuestion.objects.create(
                    exam=exam,
                    question_bank=None,
                    prompt=question['prompt'].strip(),
                    question_type=question['question_type'],
                    mark_value=question.get('mark_value', 1.0),
                    explanation=question.get('explanation', '').strip(),
                    topic=question.get('topic', '').strip(),
                    difficulty=question['difficulty'],
                    order=current_max_order + idx + 1,
                    is_active=True
                )

                for cidx, choice in enumerate(question.get('choices', [])):
                    CBTChoice.objects.create(
                        question=qobj,
                        text=choice['text'].strip(),
                        is_correct=bool(choice.get('is_correct', False)),
                        order=cidx
                    )

                created_question_ids.append(qobj.id)
    except Exception as exc:
        AIRequestMetric.objects.create(
            user=request.user,
            exam=exam,
            request_type=AIRequestMetric.REQUEST_TYPE_GENERATE_AI,
            status=AIRequestMetric.STATUS_FAILURE,
            error_code=exc.__class__.__name__,
        )
        return JsonResponse({'error': f'Failed to save generated questions: {str(exc)}'}, status=500)

    # record success metric
    try:
        AIRequestMetric.objects.create(
            user=request.user,
            exam=exam,
            request_type=AIRequestMetric.REQUEST_TYPE_GENERATE_AI,
            status=AIRequestMetric.STATUS_SUCCESS,
            latency_ms=0,
            token_usage=None,
        )
    except Exception:
        pass

    return JsonResponse({'status': 'ok', 'created_question_ids': created_question_ids})


def _normalize_text(value):
    return value.strip() if isinstance(value, str) else ''


def _parse_json_questions(text):
    try:
        data = json.loads(text)
    except json.JSONDecodeError as exc:
        raise ValueError(f'Invalid JSON content: {str(exc)}')

    if isinstance(data, dict):
        if 'questions' in data and isinstance(data['questions'], list):
            return data['questions']
        if 'items' in data and isinstance(data['items'], list):
            return data['items']
        return [data]
    if not isinstance(data, list):
        raise ValueError('JSON file must contain a list of questions or a single question object.')
    return data


def _parse_csv_questions(text):
    reader = csv.DictReader(StringIO(text))
    if not reader.fieldnames:
        raise ValueError('CSV must contain a header row.')

    questions = []
    for idx, row in enumerate(reader):
        record = {key.strip().lower(): (_normalize_text(value) if value is not None else '') for key, value in row.items()}
        raw_choices = record.get('choices', '')
        raw_correct = record.get('correct_answers', record.get('correct_answer', ''))
        correct_values = [item.strip() for item in re.split(r'[;|,]', raw_correct) if item.strip()]
        choices = []
        for choice_text in re.split(r'[;|,]', raw_choices):
            choice_text = choice_text.strip()
            if choice_text:
                choices.append({
                    'text': choice_text,
                    'is_correct': choice_text in correct_values,
                })

        questions.append({
            'prompt': record.get('prompt', ''),
            'question_type': record.get('question_type', '') or CBTQuestion.MCQ,
            'mark_value': float(record.get('mark_value', '1') or 1),
            'difficulty': record.get('difficulty', CBTQuestion.DIFFICULTY_MEDIUM),
            'topic': record.get('topic', ''),
            'explanation': record.get('explanation', ''),
            'tags': record.get('tags', ''),
            'choices': choices,
            'order': int(record.get('order', idx) or idx),
        })
    return questions


def _extract_json_from_text(text):
    content = text.strip()
    if content.startswith('{') or content.startswith('['):
        return _parse_json_questions(content)

    match = re.search(r'([\[{].*[\]}])', content, re.S)
    if match:
        return _parse_json_questions(match.group(1))

    return None


def _parse_plain_text_questions(text):
    sections = [section.strip() for section in re.split(r'\n\s*\n+', text.strip()) if section.strip()]
    questions = []
    for section in sections:
        if not section:
            continue
        json_candidate = _extract_json_from_text(section)
        if json_candidate is not None:
            questions.extend(json_candidate)
            continue

        data = {
            'prompt': '',
            'question_type': CBTQuestion.MCQ,
            'mark_value': 1.0,
            'difficulty': CBTQuestion.DIFFICULTY_MEDIUM,
            'topic': '',
            'explanation': '',
            'tags': '',
            'choices': [],
            'order': 0,
        }
        current_choices = []
        for line in section.splitlines():
            content = line.strip()
            lower = content.lower()
            if lower.startswith('prompt:') or lower.startswith('question:'):
                data['prompt'] = content.split(':', 1)[1].strip()
                continue
            if lower.startswith('type:'):
                data['question_type'] = content.split(':', 1)[1].strip() or CBTQuestion.MCQ
                continue
            if lower.startswith('marks:') or lower.startswith('mark_value:'):
                try:
                    data['mark_value'] = float(content.split(':', 1)[1].strip() or 1)
                except ValueError:
                    data['mark_value'] = 1.0
                continue
            if lower.startswith('difficulty:'):
                data['difficulty'] = content.split(':', 1)[1].strip() or CBTQuestion.DIFFICULTY_MEDIUM
                continue
            if lower.startswith('topic:'):
                data['topic'] = content.split(':', 1)[1].strip()
                continue
            if lower.startswith('explanation:'):
                data['explanation'] = content.split(':', 1)[1].strip()
                continue
            if lower.startswith('tags:'):
                data['tags'] = content.split(':', 1)[1].strip()
                continue
            if lower.startswith('choices:') or lower.startswith('options:'):
                current_choices = []
                continue
            if content.startswith('-') or content.startswith('*'):
                choice_text = content[1:].strip()
                is_correct = choice_text.startswith('*') or choice_text.startswith('✓')
                if is_correct:
                    choice_text = choice_text.lstrip('*✓').strip()
                current_choices.append({'text': choice_text, 'is_correct': is_correct})
                continue
            if current_choices and content and not lower.startswith('correct') and not lower.startswith('answer') and not lower.startswith('correct answers:'):
                current_choices.append({'text': content, 'is_correct': False})

        if current_choices:
            data['choices'] = current_choices
        else:
            data['question_type'] = CBTQuestion.SHORT_ANSWER
            data['choices'] = []

        questions.append(data)
    return questions


def _load_questions_from_docx(file_obj):
    try:
        from docx import Document as DocxDocument
    except ImportError:
        raise ValueError('DOCX import requires the python-docx library.')

    document = DocxDocument(file_obj)
    text = '\n\n'.join([paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()])
    if not text:
        raise ValueError('Word document is empty.')

    json_questions = _extract_json_from_text(text)
    if json_questions is not None:
        return json_questions
    return _parse_plain_text_questions(text)


def _load_questions_from_xlsx(file_obj):
    try:
        import openpyxl
    except ImportError:
        raise ValueError('Excel import requires the openpyxl library.')

    workbook = openpyxl.load_workbook(filename=BytesIO(file_obj.read()), read_only=True, data_only=True)
    sheet = workbook.active
    headers = [cell.strip().lower() if isinstance(cell, str) else '' for cell in next(sheet.iter_rows(values_only=True))]
    questions = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        record = {headers[idx]: (_normalize_text(value) if isinstance(value, str) else (str(value) if value is not None else '')) for idx, value in enumerate(row) if idx < len(headers)}
        raw_choices = record.get('choices', '')
        raw_correct = record.get('correct_answers', record.get('correct_answer', ''))
        correct_values = [item.strip() for item in re.split(r'[;|,]', raw_correct) if item.strip()]
        choices = []
        for choice_text in re.split(r'[;|,]', raw_choices):
            choice_text = choice_text.strip()
            if choice_text:
                choices.append({
                    'text': choice_text,
                    'is_correct': choice_text in correct_values,
                })
        questions.append({
            'prompt': record.get('prompt', ''),
            'question_type': record.get('question_type', '') or CBTQuestion.MCQ,
            'mark_value': float(record.get('mark_value', '1') or 1),
            'difficulty': record.get('difficulty', CBTQuestion.DIFFICULTY_MEDIUM),
            'topic': record.get('topic', ''),
            'explanation': record.get('explanation', ''),
            'tags': record.get('tags', ''),
            'choices': choices,
            'order': int(record.get('order', 0) or 0),
        })
    return questions


def _serialize_question(q):
    return {
        'prompt': q.prompt,
        'question_type': q.question_type,
        'mark_value': float(q.mark_value),
        'difficulty': q.difficulty,
        'topic': q.topic,
        'explanation': q.explanation,
        'tags': q.tags,
        'choices': [{'text': c.text, 'is_correct': c.is_correct, 'order': c.order} for c in q.choices.all().order_by('order')],
        'order': q.order,
        'is_active': q.is_active,
    }


def _build_export_filename(exam, suffix):
    name = slugify(exam.name or f'exam-{exam.pk}') or f'exam-{exam.pk}'
    return f'{name}-{suffix}'


@login_required
def import_exam_questions(request, exam_pk):
    if request.method != 'POST':
        return JsonResponse({'error': 'POST required'}, status=405)

    exam = get_object_or_404(CBTExam, pk=exam_pk, created_by=request.user)
    importer = QuestionImporter(user=request.user, target_exam=exam)

    if request.content_type and 'application/json' in request.content_type:
        try:
            payload = json.loads(request.body.decode('utf-8'))
        except Exception:
            return JsonResponse({'error': 'Invalid JSON payload.'}, status=400)

        mode = payload.get('mode', 'save')
        import_target = payload.get('import_target', 'exam')
        bank_id = payload.get('bank_id')
        selected_rows = payload.get('rows')
        if not isinstance(selected_rows, list):
            return JsonResponse({'error': 'rows must be a list.'}, status=400)

        if mode == 'preview':
            rows = []
            for idx, row in enumerate(selected_rows, start=1):
                question_row = importer._normalize_row(row, idx)
                rows.append(question_row)
            preview = ImportPreviewSerializer.serialize(rows)
            return JsonResponse({'status': 'ok', 'preview': preview})

        if mode == 'save':
            if import_target in ('bank', 'both') and not bank_id:
                return JsonResponse({'error': 'bank_id is required when importing into a question bank.'}, status=400)
            bank = None
            if bank_id:
                bank = get_object_or_404(QuestionBank, pk=bank_id, created_by=request.user)
            result = importer.save(selected_rows, import_target=import_target, bank=bank)
            response_data = {'status': 'ok', **result}
            return JsonResponse(response_data)

        return JsonResponse({'error': 'Invalid import mode.'}, status=400)

    upload = request.FILES.get('file')
    if not upload:
        return JsonResponse({'error': 'No file uploaded.'}, status=400)

    save_direct = request.POST.get('save_direct') == '1' or request.POST.get('mode') == 'save'
    preview_mode = request.POST.get('mode') == 'preview'
    import_target = request.POST.get('import_target', 'exam')
    bank_id = request.POST.get('bank_id')

    try:
        rows = importer.parse_file(upload)
    except QuestionImportError as exc:
        return JsonResponse({'error': str(exc)}, status=400)
    except Exception as exc:
        return JsonResponse({'error': f'Failed to parse import file: {str(exc)}'}, status=500)

    if preview_mode and not save_direct:
        preview = ImportPreviewSerializer.serialize(rows)
        return JsonResponse({'status': 'ok', 'preview': preview})

    bank = None
    if import_target in ('bank', 'both'):
        if not bank_id:
            return JsonResponse({'error': 'bank_id is required when importing into a question bank.'}, status=400)
        bank = get_object_or_404(QuestionBank, pk=bank_id, created_by=request.user)

    try:
        row_dicts = [row.normalized for row in rows if row.valid]
        result = importer.save(row_dicts, import_target=import_target, bank=bank)
        response_data = {'status': 'ok', 'preview': ImportPreviewSerializer.serialize(rows), **result}
        return JsonResponse(response_data)
    except QuestionImportError as exc:
        return JsonResponse({'error': str(exc)}, status=400)
    except Exception as exc:
        return JsonResponse({'error': f'Failed to import questions: {str(exc)}'}, status=500)


@login_required
def download_question_template(request, exam_pk):
    exam = get_object_or_404(CBTExam, pk=exam_pk, created_by=request.user)
    fmt = (request.GET.get('format') or 'json').lower()
    template_data = [
        {
            'prompt': 'What is 2 + 2?',
            'question_type': 'mcq',
            'mark_value': 1,
            'difficulty': 'medium',
            'topic': 'Arithmetic',
            'explanation': 'Provide the numeric answer.',
            'tags': 'math,addition',
            'choices': [
                {'text': '4', 'is_correct': True},
                {'text': '3', 'is_correct': False},
                {'text': '5', 'is_correct': False},
            ],
        }
    ]
    filename_base = _build_export_filename(exam, 'template')
    if fmt == 'csv':
        output = StringIO()
        writer = csv.writer(output)
        writer.writerow(['prompt', 'question_type', 'mark_value', 'difficulty', 'topic', 'explanation', 'tags', 'choices', 'correct_answers'])
        writer.writerow([
            template_data[0]['prompt'],
            template_data[0]['question_type'],
            template_data[0]['mark_value'],
            template_data[0]['difficulty'],
            template_data[0]['topic'],
            template_data[0]['explanation'],
            template_data[0]['tags'],
            ';'.join([choice['text'] for choice in template_data[0]['choices']]),
            ';'.join([choice['text'] for choice in template_data[0]['choices'] if choice['is_correct']]),
        ])
        response = HttpResponse(output.getvalue(), content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="{slugify(exam.name) or "exam"}-question-template.csv"'
        return response
    if fmt in ('txt', 'text'):
        lines = [
            'Prompt: What is 2 + 2?',
            'Type: mcq',
            'Marks: 1',
            'Difficulty: medium',
            'Topic: Arithmetic',
            'Explanation: Provide the numeric answer.',
            'Tags: math, addition',
            'Choices:',
            '- 4',
            '- 3',
            '- 5',
            'Correct Answers: 4',
        ]
        response = HttpResponse('\n'.join(lines), content_type='text/plain')
        response['Content-Disposition'] = f'attachment; filename="{slugify(exam.name) or "exam"}-question-template.txt"'
        return response
    if fmt == 'json':
        response = HttpResponse(json.dumps(template_data, indent=2), content_type='application/json')
        response['Content-Disposition'] = f'attachment; filename="{slugify(exam.name) or "exam"}-question-template.json"'
        return response

    return JsonResponse({'error': 'Unsupported template format.'}, status=400)


@login_required
def export_exam_questions(request, exam_pk, format):
    exam = get_object_or_404(CBTExam, pk=exam_pk, created_by=request.user)
    questions = [_serialize_question(q) for q in exam.questions.prefetch_related('choices').order_by('order')]
    filename_base = slugify(exam.name) or f'exam-{exam.pk}'
    if format == 'json':
        response = HttpResponse(json.dumps(questions, indent=2), content_type='application/json')
        response['Content-Disposition'] = f'attachment; filename="{filename_base}-questions.json"'
        return response
    if format == 'csv':
        output = StringIO()
        writer = csv.writer(output)
        writer.writerow(['prompt', 'question_type', 'mark_value', 'difficulty', 'topic', 'explanation', 'tags', 'order', 'choices', 'correct_answers', 'is_active'])
        for question in questions:
            choice_texts = ';'.join([choice['text'] for choice in question['choices']])
            correct_texts = ';'.join([choice['text'] for choice in question['choices'] if choice['is_correct']])
            writer.writerow([
                question['prompt'],
                question['question_type'],
                question['mark_value'],
                question['difficulty'],
                question['topic'],
                question['explanation'],
                question['tags'],
                question['order'],
                choice_texts,
                correct_texts,
                question['is_active'],
            ])
        response = HttpResponse(output.getvalue(), content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="{filename_base}-questions.csv"'
        return response
    if format in ('docx', 'word'):
        document = Document()
        document.add_heading(f'Questions for {exam.name}', level=1)
        for idx, question in enumerate(questions, start=1):
            document.add_paragraph(f'{idx}. {question["prompt"]}', style='List Number')
            document.add_paragraph(f'Type: {question["question_type"]}')
            document.add_paragraph(f'Marks: {question["mark_value"]}')
            document.add_paragraph(f'Difficulty: {question["difficulty"]}')
            if question['topic']:
                document.add_paragraph(f'Topic: {question["topic"]}')
            if question['explanation']:
                document.add_paragraph(f'Explanation: {question["explanation"]}')
            if question['tags']:
                document.add_paragraph(f'Tags: {question["tags"]}')
            for choice in question['choices']:
                item = document.add_paragraph(style='List Bullet')
                item.add_run(choice['text'])
                if choice['is_correct']:
                    item.add_run(' (correct)')
        output = BytesIO()
        document.save(output)
        output.seek(0)
        response = FileResponse(output, content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        response['Content-Disposition'] = f'attachment; filename="{filename_base}-questions.docx"'
        return response
    if format == 'pdf':
        html_lines = ['<html><head><meta charset="utf-8"><title>Question Export</title></head><body>']
        html_lines.append(f'<h1>Questions for {exam.name}</h1>')
        for idx, question in enumerate(questions, start=1):
            html_lines.append(f'<h2>{idx}. {question["prompt"]}</h2>')
            html_lines.append(f'<p><strong>Type:</strong> {question["question_type"]}</p>')
            html_lines.append(f'<p><strong>Marks:</strong> {question["mark_value"]}</p>')
            if question['topic']:
                html_lines.append(f'<p><strong>Topic:</strong> {question["topic"]}</p>')
            if question['explanation']:
                html_lines.append(f'<p><strong>Explanation:</strong> {question["explanation"]}</p>')
            if question['tags']:
                html_lines.append(f'<p><strong>Tags:</strong> {question["tags"]}</p>')
            html_lines.append('<ul>')
            for choice in question['choices']:
                label = ' (correct)' if choice['is_correct'] else ''
                html_lines.append(f'<li>{choice["text"]}{label}</li>')
            html_lines.append('</ul>')
        html_lines.append('</body></html>')
        html_string = ''.join(html_lines)
        try:
            from weasyprint import HTML
            pdf = HTML(string=html_string, base_url=request.build_absolute_uri('/')).write_pdf()
            response = HttpResponse(pdf, content_type='application/pdf')
            response['Content-Disposition'] = f'attachment; filename="{filename_base}-questions.pdf"'
            return response
        except ImportError:
            return JsonResponse({'error': 'PDF export requires WeasyPrint.'}, status=501)
        except Exception as exc:
            return JsonResponse({'error': f'Failed to generate PDF: {str(exc)}'}, status=500)
    if format in ('xlsx', 'excel'):
        try:
            import openpyxl
            from openpyxl import Workbook
        except ImportError:
            return JsonResponse({'error': 'Excel export requires the openpyxl library.'}, status=501)
        wb = Workbook()
        ws = wb.active
        ws.title = 'Questions'
        ws.append(['prompt', 'question_type', 'mark_value', 'difficulty', 'topic', 'explanation', 'tags', 'order', 'choices', 'correct_answers', 'is_active'])
        for question in questions:
            choice_texts = ';'.join([choice['text'] for choice in question['choices']])
            correct_texts = ';'.join([choice['text'] for choice in question['choices'] if choice['is_correct']])
            ws.append([
                question['prompt'],
                question['question_type'],
                question['mark_value'],
                question['difficulty'],
                question['topic'],
                question['explanation'],
                question['tags'],
                question['order'],
                choice_texts,
                correct_texts,
                question['is_active'],
            ])
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        response = FileResponse(output, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="{filename_base}-questions.xlsx"'
        return response
    if format in ('backup', 'zip'):
        archive = BytesIO()
        with zipfile.ZipFile(archive, 'w', zipfile.ZIP_DEFLATED) as zip_file:
            zip_file.writestr(f'{filename_base}-questions.json', json.dumps(questions, indent=2))
            csv_output = StringIO()
            csv_writer = csv.writer(csv_output)
            csv_writer.writerow(['prompt', 'question_type', 'mark_value', 'difficulty', 'topic', 'explanation', 'tags', 'order', 'choices', 'correct_answers', 'is_active'])
            for question in questions:
                choice_texts = ';'.join([choice['text'] for choice in question['choices']])
                correct_texts = ';'.join([choice['text'] for choice in question['choices'] if choice['is_correct']])
                csv_writer.writerow([
                    question['prompt'],
                    question['question_type'],
                    question['mark_value'],
                    question['difficulty'],
                    question['topic'],
                    question['explanation'],
                    question['tags'],
                    question['order'],
                    choice_texts,
                    correct_texts,
                    question['is_active'],
                ])
            zip_file.writestr(f'{filename_base}-questions.csv', csv_output.getvalue())
        archive.seek(0)
        response = FileResponse(archive, content_type='application/zip')
        response['Content-Disposition'] = f'attachment; filename="{filename_base}-questions-backup.zip"'
        return response

    return JsonResponse({'error': 'Unsupported export format.'}, status=400)


    return JsonResponse({'error': 'Unsupported export format.'}, status=400)


@method_decorator(login_required, name='dispatch')
class StudentCBTDashboardView(LoginRequiredMixin, UserPassesTestMixin, TemplateView):
    template_name = 'cbt/student_dashboard.html'

    def test_func(self):
        return is_cbt_student(self.request.user)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        now = timezone.now()
        # include attempts either owned by the user or tied to the current session
        session_key = _ensure_session(self.request)
        context['active_attempts'] = CBTStudentAttempt.objects.filter(
            (Q(student=self.request.user) | Q(session_key=session_key)),
            is_submitted=False
        ).select_related('exam').order_by('-started_at')
        # provide explicit counts to make templates defensive and straightforward
        context['active_attempts_count'] = context['active_attempts'].count()
        context['recent_cbt_results'] = CBTStudentAttempt.objects.filter(
            (Q(student=self.request.user) | Q(session_key=session_key)),
            is_submitted=True
        ).select_related('exam').order_by('-completed_at')[:5]
        context['recent_cbt_results_count'] = CBTStudentAttempt.objects.filter((Q(student=self.request.user) | Q(session_key=session_key)), is_submitted=True).count()
        student_class = getattr(getattr(self.request.user, 'student_profile', None), 'student_class', None)
        if student_class:
            context['upcoming_cbt_exams'] = CBTExam.objects.filter(
                exam_mode=CBTExam.REAL,
                is_active=True,
                is_published=True,
                start_datetime__gt=now,
                school_class=student_class
            ).order_by('start_datetime')
            context['upcoming_cbt_exams_count'] = context['upcoming_cbt_exams'].count()
        else:
            context['upcoming_cbt_exams'] = CBTExam.objects.none()
            context['upcoming_cbt_exams_count'] = 0
        return context


@method_decorator(login_required, name='dispatch')
class StudentCBTPracticeListView(LoginRequiredMixin, UserPassesTestMixin, PracticeExamListView):
    template_name = 'cbt/student_practice_list.html'

    def test_func(self):
        return is_cbt_student(self.request.user)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        student_class = getattr(getattr(self.request.user, 'student_profile', None), 'student_class', None)
        if student_class:
            context['practice_exams'] = context['practice_exams'].filter(school_class=student_class)
        return context


@method_decorator(login_required, name='dispatch')
class StudentCBTAttemptListView(LoginRequiredMixin, UserPassesTestMixin, TemplateView):
    template_name = 'cbt/student_attempt_list.html'

    def test_func(self):
        return is_cbt_student(self.request.user)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        session_key = _ensure_session(self.request)
        context['attempts'] = CBTStudentAttempt.objects.filter(
            (Q(student=self.request.user) | Q(session_key=session_key))
        ).select_related('exam').order_by('-started_at')
        return context


@method_decorator(login_required, name='dispatch')
class StudentCBTResultListView(LoginRequiredMixin, UserPassesTestMixin, TemplateView):
    template_name = 'cbt/student_result_list.html'

    def test_func(self):
        return is_cbt_student(self.request.user)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        session_key = _ensure_session(self.request)
        context['attempts'] = CBTStudentAttempt.objects.filter(
            (Q(student=self.request.user) | Q(session_key=session_key)),
            is_submitted=True
        ).select_related('exam').order_by('-completed_at')
        return context


@method_decorator(login_required, name='dispatch')
class TeacherCBTDashboardView(LoginRequiredMixin, UserPassesTestMixin, TemplateView):
    template_name = 'cbt/teacher_dashboard.html'

    def test_func(self):
        return is_cbt_teacher(self.request.user)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        exams = CBTExam.objects.filter(created_by=self.request.user)
        attempts = CBTStudentAttempt.objects.filter(exam__created_by=self.request.user)
        context['cbt_exams_created'] = exams.count()
        context['cbt_active_exams'] = exams.filter(is_active=True, is_published=True).count()
        context['cbt_attempts'] = attempts.count()
        context['cbt_question_banks'] = QuestionBank.objects.filter(created_by=self.request.user).count()
        context['cbt_question_bank_questions'] = CBTQuestion.objects.filter(question_bank__created_by=self.request.user).count()
        context['cbt_recent_question_banks'] = QuestionBank.objects.filter(created_by=self.request.user).order_by('-created_at')[:5]
        context['cbt_recent_attempts'] = attempts.select_related('student', 'exam').order_by('-started_at')[:5]
        completed_attempts = attempts.filter(is_submitted=True)
        context['cbt_avg_score'] = completed_attempts.aggregate(avg_score=Avg('score'))['avg_score'] or 0
        return context


@method_decorator(login_required, name='dispatch')
class TeacherCBTExamListView(CBTExamListView):
    template_name = 'cbt/teacher_exam_list.html'

    def get_queryset(self):
        return super().get_queryset().filter(created_by=self.request.user)


@method_decorator(login_required, name='dispatch')
class TeacherCBTExamCreateView(CBTExamCreateView):
    success_url = reverse_lazy('teacher_cbt:manage')


@method_decorator(login_required, name='dispatch')
class TeacherCBTExamUpdateView(CBTExamUpdateView):
    success_url = reverse_lazy('teacher_cbt:manage')


@method_decorator(login_required, name='dispatch')
class TeacherCBTAttemptListView(LoginRequiredMixin, UserPassesTestMixin, ListView):
    model = CBTStudentAttempt
    template_name = 'cbt/teacher_attempt_list.html'
    context_object_name = 'attempts'
    paginate_by = 20

    def test_func(self):
        return is_cbt_teacher(self.request.user)

    def get_queryset(self):
        return CBTStudentAttempt.objects.filter(
            exam__created_by=self.request.user
        ).select_related('exam', 'student').annotate(
            integrity_count=Count('integrity_events')
        ).order_by('-started_at')


@method_decorator(login_required, name='dispatch')
class TeacherCBTAnalyticsView(LoginRequiredMixin, UserPassesTestMixin, TemplateView):
    template_name = 'cbt/teacher_analytics.html'

    def test_func(self):
        return is_cbt_teacher(self.request.user)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        exams = CBTExam.objects.filter(created_by=self.request.user)
        attempts = CBTStudentAttempt.objects.filter(exam__created_by=self.request.user, is_submitted=True)
        context['cbt_exams_created'] = exams.count()
        context['cbt_active_exams'] = exams.filter(is_active=True, is_published=True).count()
        context['cbt_attempts'] = attempts.count()
        context['cbt_avg_score'] = attempts.aggregate(avg_score=Avg('score'))['avg_score'] or 0
        today = timezone.localdate()
        ai_metrics = AIRequestMetric.objects.filter(exam__created_by=self.request.user)
        context['ai_metrics_summary'] = ai_metrics.values('request_type', 'status').annotate(count=Count('id'))
        context['ai_metrics_today'] = ai_metrics.filter(date=today).values('status').annotate(count=Count('id'))
        return context


@method_decorator(login_required, name='dispatch')
class TeacherCBTAIGeneratorView(LoginRequiredMixin, UserPassesTestMixin, TemplateView):
    template_name = 'cbt/ai_question_generator.html'

    def test_func(self):
        return (
            getattr(settings, 'AI_GENERATION_ENABLED', True)
            and CBTExam.objects.filter(
                pk=self.kwargs.get('exam_pk'),
                created_by=self.request.user,
                allow_ai_questions=True,
            ).exists()
        )

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        exam = get_object_or_404(CBTExam, pk=self.kwargs.get('exam_pk'))
        context['exam'] = exam
        context['QUESTION_TYPE_CHOICES'] = CBTQuestion.QUESTION_TYPE_CHOICES
        context['DIFFICULTY_CHOICES'] = CBTQuestion.DIFFICULTY_CHOICES
        context['ai_generation_enabled'] = getattr(settings, 'AI_GENERATION_ENABLED', True)
        return context


class ManageExamQuestionsView(LoginRequiredMixin, UserPassesTestMixin, TemplateView):
    """Bulk question builder and manager for a specific exam"""
    template_name = 'cbt/manage_exam_questions.html'

    def test_func(self):
        exam = get_object_or_404(CBTExam, pk=self.kwargs.get('exam_pk'))
        return exam.created_by == self.request.user

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        exam = get_object_or_404(CBTExam, pk=self.kwargs.get('exam_pk'))
        # existing questions for the exam
        context['exam'] = exam
        context['questions'] = exam.questions.all().prefetch_related('choices')
        # include question bank options
        context['question_banks'] = QuestionBank.objects.filter(created_by=self.request.user)
        # provide choice lists for template
        context['QUESTION_TYPE_CHOICES'] = CBTQuestion.QUESTION_TYPE_CHOICES
        context['DIFFICULTY_CHOICES'] = CBTQuestion.DIFFICULTY_CHOICES
        return context

    def post(self, request, *args, **kwargs):
        exam = get_object_or_404(CBTExam, pk=self.kwargs.get('exam_pk'))
        if exam.created_by != request.user:
            return HttpResponseForbidden()

        payload = {}
        if request.content_type and 'application/json' in request.content_type:
            try:
                payload = json.loads(request.body.decode())
            except Exception as e:
                return JsonResponse({'error': f'Invalid JSON payload: {str(e)}'}, status=400)
        else:
            try:
                payload_str = request.POST.get('payload', '{}')
                if not payload_str or payload_str == '{}':
                    return JsonResponse({'error': 'No payload provided in request'}, status=400)
                payload = json.loads(payload_str)
            except Exception as e:
                return JsonResponse({'error': f'Invalid payload: {str(e)}'}, status=400)

        questions = payload.get('questions', [])
        bank_import = payload.get('bank_import')
        deleted_ids = payload.get('deleted_question_ids', [])

        created = []
        updated = []
        deleted = []

        from django.db import transaction

        try:
            with transaction.atomic():
                # delete removed questions
                if deleted_ids:
                    for qid in deleted_ids:
                        qobj = CBTQuestion.objects.filter(pk=qid, exam=exam).first()
                        if qobj:
                            qobj.delete()
                            deleted.append(qid)

                for idx, q in enumerate(questions):
                    existing_id = q.get('existing_id')
                    if existing_id:
                        qobj = CBTQuestion.objects.filter(pk=existing_id, exam=exam).first()
                    else:
                        qobj = None

                    if qobj:
                        qobj.prompt = q.get('prompt', qobj.prompt)
                        qobj.question_type = q.get('question_type', qobj.question_type)
                        qobj.mark_value = q.get('mark_value', qobj.mark_value)
                        qobj.explanation = q.get('explanation', qobj.explanation)
                        qobj.topic = q.get('topic', qobj.topic)
                        qobj.difficulty = q.get('difficulty', qobj.difficulty)
                        qobj.order = q.get('order', qobj.order)
                        qobj.is_active = q.get('is_active', True)
                        qobj.save()
                        updated.append(qobj.id)
                        qobj.choices.all().delete()
                    else:
                        qobj = CBTQuestion.objects.create(
                            exam=exam,
                            question_bank=None,
                            prompt=q.get('prompt', ''),
                            question_type=q.get('question_type', CBTQuestion.MCQ),
                            mark_value=q.get('mark_value', 1.0),
                            explanation=q.get('explanation', ''),
                            topic=q.get('topic', ''),
                            difficulty=q.get('difficulty', CBTQuestion.DIFFICULTY_MEDIUM),
                            order=q.get('order', idx),
                            is_active=q.get('is_active', True)
                        )
                        created.append(qobj.id)

                    if request.FILES:
                        image_key = f'image_{existing_id or qobj.id}'
                        if image_key in request.FILES:
                            qobj.image = request.FILES[image_key]
                            qobj.save()

                    for cidx, ch in enumerate(q.get('choices', [])):
                        CBTChoice.objects.create(
                            question=qobj,
                            text=ch.get('text', ''),
                            is_correct=bool(ch.get('is_correct', False)),
                            order=ch.get('order', cidx)
                        )

                if bank_import:
                    bank_id = bank_import.get('bank_id')
                    selected_ids = bank_import.get('question_ids', [])
                    qb = get_object_or_404(QuestionBank, pk=bank_id, created_by=request.user)
                    for qid in selected_ids:
                        src = get_object_or_404(CBTQuestion, pk=qid, question_bank=qb)
                        nq = CBTQuestion.objects.create(
                            exam=exam,
                            question_bank=qb,
                            prompt=src.prompt,
                            question_type=src.question_type,
                            mark_value=src.mark_value,
                            explanation=src.explanation,
                            topic=src.topic,
                            difficulty=src.difficulty,
                            order=src.order,
                            is_active=src.is_active
                        )
                        for c in src.choices.all():
                            CBTChoice.objects.create(
                                question=nq,
                                text=c.text,
                                is_correct=c.is_correct,
                                order=c.order
                            )
                        created.append(nq.id)

        except Exception as e:
            return JsonResponse({'error': f'Save failed: {str(e)}'}, status=500)

        return JsonResponse({'status': 'ok', 'created': created, 'updated': updated, 'deleted': deleted})
